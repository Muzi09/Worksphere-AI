import io
import os
import docx
import fitz  # PyMuPDF
from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, status
from app.core.deps import get_current_user
from app.models.user import User

router = APIRouter()

def extract_text_from_pdf(file_bytes: bytes) -> str:
    pdf_doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_text = []
    for page in pdf_doc:
        text = page.get_text("text")
        if text.strip():
            pages_text.append(text.strip())
    pdf_doc.close()
    return "\n\n".join(pages_text)

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            extracted_text.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                extracted_text.append(" | ".join(row_text))
    return "\n".join(extracted_text)

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        try:
            return extract_text_from_pdf(file_bytes)
        except Exception as e:
            print(f"Error parsing PDF with PyMuPDF: {e}")

    if ext == ".docx":
        try:
            return extract_text_from_docx(file_bytes)
        except Exception as e:
            print(f"Error parsing DOCX: {e}")

    # Fallback attempt 1: PDF via PyMuPDF
    try:
        pdf_text = extract_text_from_pdf(file_bytes)
        if pdf_text:
            return pdf_text
    except Exception:
        pass

    # Fallback attempt 2: DOCX via python-docx
    try:
        docx_text = extract_text_from_docx(file_bytes)
        if docx_text:
            return docx_text
    except Exception:
        pass

    # Fallback attempt 3: Plain text decoding (UTF-8 / Latin-1)
    try:
        return file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        return file_bytes.decode('latin-1', errors='ignore')

@router.post("/file-upload")
async def upload_file(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user)
):
    try:
        contents = await file.read()
        file_size = len(contents)

        extracted_text = extract_text_from_file(contents, file.filename)
        

        return {
            "status": "success",
            "message": f"File '{file.filename}' uploaded and converted to text successfully",
            "filename": file.filename,
            "content_type": file.content_type,
            "size": file_size,
            "text": extracted_text
        }
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to process file upload: {str(e)}"
        )
